from __future__ import annotations

import csv
import json
import os
import re
import urllib.request
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from app.services.jarvis_internet_tools import internet_search, internet_research

ROOT = Path.cwd()
ART_DIR = ROOT / "jarvis_stage3_artifacts" / "telegram_tables"
ART_DIR.mkdir(parents=True, exist_ok=True)


def _safe_name(text: str) -> str:
    cleaned = "".join(ch if ch.isalnum() else "_" for ch in text[:60])
    return cleaned.strip("_") or "jarvis_table"


def _send_telegram_document(file_path: str, caption: str = "") -> Dict[str, Any]:
    token = os.getenv("TELEGRAM_BOT_TOKEN", "").strip()
    chat_id = os.getenv("TELEGRAM_ALLOWED_CHAT_ID", "").strip()

    if not token:
        return {"ok": False, "error": "TELEGRAM_BOT_TOKEN is empty", "file_path": file_path}
    if not chat_id:
        return {"ok": False, "error": "TELEGRAM_ALLOWED_CHAT_ID is empty", "file_path": file_path}

    boundary = "----JarvisBoundary" + datetime.utcnow().strftime("%Y%m%d%H%M%S%f")
    url = f"https://api.telegram.org/bot{token}/sendDocument"

    file_name = Path(file_path).name
    file_bytes = Path(file_path).read_bytes()

    parts = []
    def add_field(name: str, value: str):
        parts.append(f"--{boundary}\r\n".encode())
        parts.append(f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode())
        parts.append(str(value).encode("utf-8"))
        parts.append(b"\r\n")

    add_field("chat_id", chat_id)
    if caption:
        add_field("caption", caption[:1000])

    parts.append(f"--{boundary}\r\n".encode())
    parts.append(f'Content-Disposition: form-data; name="document"; filename="{file_name}"\r\n'.encode())
    parts.append(b"Content-Type: application/octet-stream\r\n\r\n")
    parts.append(file_bytes)
    parts.append(b"\r\n")
    parts.append(f"--{boundary}--\r\n".encode())

    body = b"".join(parts)

    req = urllib.request.Request(
        url,
        data=body,
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
        method="POST",
    )

    try:
        with urllib.request.urlopen(req, timeout=120) as resp:
            raw = resp.read().decode("utf-8", errors="replace")
            data = json.loads(raw)
            return {"ok": bool(data.get("ok")), "telegram": data, "file_path": file_path}
    except Exception as e:
        return {"ok": False, "error": str(e), "file_path": file_path}


def _post_json(url: str, payload: Dict[str, Any], headers: Dict[str, str], timeout: int = 30) -> Dict[str, Any]:
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    req = urllib.request.Request(url, data=body, headers={**headers, "Content-Type": "application/json"}, method="POST")
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return json.loads(r.read().decode("utf-8", errors="replace"))
    except Exception as e:
        return {"error": str(e)}


def _llm_extract_table(
    query: str,
    search_results: List[Dict[str, Any]],
    research_answer: str,
) -> Optional[Dict[str, Any]]:
    """Call LLM to produce topic-appropriate columns and rows as JSON.
    Returns {"columns": [...], "rows": [[...], ...]} or None on failure.
    """
    items_text = "\n".join(
        f"{i+1}. {r.get('title', '')} — {r.get('url', '')}\n   {r.get('content', '')[:300]}"
        for i, r in enumerate(search_results[:10])
    )
    prompt = (
        f"Topic: {query}\n\n"
        f"Search results:\n{items_text}\n\n"
        f"Research summary:\n{research_answer[:2000]}\n\n"
        "Task: Create a structured comparison table with 8-12 rows.\n"
        "Choose columns that are most useful for this topic (e.g. for AI services: "
        "Service, Category, Pricing, Free Tier, API Available, Strengths, Limitations, Website).\n"
        "Rows should contain REAL data extracted from the search results and research.\n"
        "Return ONLY valid JSON, no markdown, no explanation:\n"
        '{"columns": ["Col1", "Col2", ...], "rows": [["val1", "val2", ...], ...]}'
    )

    openai_key = os.getenv("OPENAI_API_KEY", "").strip()
    if openai_key:
        resp = _post_json(
            "https://api.openai.com/v1/chat/completions",
            {
                "model": "gpt-4o-mini",
                "messages": [{"role": "user", "content": prompt}],
                "temperature": 0.2,
                "max_tokens": 2000,
            },
            {"Authorization": f"Bearer {openai_key}"},
            timeout=40,
        )
        text = ((resp.get("choices") or [{}])[0].get("message") or {}).get("content", "")
        if text:
            return _parse_table_json(text)

    anthropic_key = os.getenv("ANTHROPIC_API_KEY", "").strip()
    if anthropic_key:
        resp = _post_json(
            "https://api.anthropic.com/v1/messages",
            {
                "model": "claude-haiku-4-5-20251001",
                "max_tokens": 2000,
                "messages": [{"role": "user", "content": prompt}],
            },
            {
                "x-api-key": anthropic_key,
                "anthropic-version": "2023-06-01",
            },
            timeout=40,
        )
        text = ((resp.get("content") or [{}])[0]).get("text", "")
        if text:
            return _parse_table_json(text)

    return None


def _parse_table_json(text: str) -> Optional[Dict[str, Any]]:
    """Extract JSON from LLM response, handling markdown code fences."""
    # Strip markdown code fences if present
    text = re.sub(r"```(?:json)?\s*", "", text).strip()
    try:
        data = json.loads(text)
        cols = data.get("columns")
        rows = data.get("rows")
        if isinstance(cols, list) and isinstance(rows, list) and cols and rows:
            return {"columns": cols, "rows": rows}
    except Exception:
        pass
    return None


def build_internet_table(query: str, max_results: int = 8, send_to_telegram: bool = True) -> Dict[str, Any]:
    ts = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    base = _safe_name(query)
    csv_path = ART_DIR / f"{base}_{ts}.csv"
    xlsx_path = ART_DIR / f"{base}_{ts}.xlsx"
    json_path = ART_DIR / f"{base}_{ts}.json"

    search = internet_search(query, max_results=max_results)
    research = internet_research(
        "Summarize this search topic and give practical recommendations: " + query
    )

    raw_results: List[Dict[str, Any]] = (((search or {}).get("data") or {}).get("results") or [])
    research_answer: str = (research or {}).get("answer", "") or ""

    # Try LLM-structured extraction first
    smart = _llm_extract_table(query, raw_results, research_answer)

    if smart:
        columns = smart["columns"]
        smart_rows = smart["rows"]
        # Pad/truncate each row to match columns length
        ncols = len(columns)
        smart_rows = [
            (list(r) + [""] * ncols)[:ncols] for r in smart_rows
        ]

        with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(columns)
            writer.writerows(smart_rows)

        rows_count = len(smart_rows)
        used_smart = True
    else:
        # Fallback: raw Tavily results
        fallback_rows = [
            [i + 1, r.get("title", ""), r.get("url", ""), r.get("score", ""), r.get("content", "")]
            for i, r in enumerate(raw_results)
        ]
        columns = ["rank", "title", "url", "score", "summary"]
        smart_rows = fallback_rows
        rows_count = len(fallback_rows)
        used_smart = False

        with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(columns)
            writer.writerows(fallback_rows)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "Internet Research"

        ws.append(columns)
        for row in smart_rows:
            ws.append(row)

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.alignment = Alignment(horizontal="center")

        col_width = max(18, 120 // max(len(columns), 1))
        for idx in range(1, len(columns) + 1):
            ws.column_dimensions[get_column_letter(idx)].width = col_width

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        wb.save(xlsx_path)
        table_path = str(xlsx_path)
    except Exception:
        table_path = str(csv_path)

    report = {
        "ok": True,
        "query": query,
        "max_results": max_results,
        "csv_path": str(csv_path),
        "xlsx_path": str(xlsx_path),
        "table_path": table_path,
        "rows_count": rows_count,
        "columns": columns,
        "used_smart_extraction": used_smart,
        "search": search,
        "research": research,
    }

    json_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    report["json_path"] = str(json_path)

    if send_to_telegram:
        caption = f"Jarvis internet table\nQuery: {query}\nRows: {rows_count}"
        report["telegram_send"] = _send_telegram_document(table_path, caption)

    return report