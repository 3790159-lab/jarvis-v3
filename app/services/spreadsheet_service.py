from __future__ import annotations

import csv
import json
import os
import shutil
import traceback
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from googleapiclient.errors import HttpError
except Exception:
    HttpError = None


BACKUP_DIR = Path("artifacts/spreadsheet_backups")
LOG_FILE = Path("artifacts/logs/spreadsheet_google_errors.log")
DEFAULT_OAUTH_TOKEN_PATH = Path("state/google_oauth_token.json")

GOOGLE_SCOPES = [
    "https://www.googleapis.com/auth/spreadsheets",
    "https://www.googleapis.com/auth/drive",
]


def _utc_now_stamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")


def _log_error(message: str) -> None:
    LOG_FILE.parent.mkdir(parents=True, exist_ok=True)
    with LOG_FILE.open("a", encoding="utf-8") as f:
        f.write(f"[{_utc_now_stamp()}] {message}\n")


def _format_exception(exc: Exception) -> str:
    parts: List[str] = []

    text = str(exc).strip()
    if text:
        parts.append(text)

    if HttpError is not None and isinstance(exc, HttpError):
        try:
            status_code = getattr(exc.resp, "status", None)
            reason = getattr(exc.resp, "reason", None)
            if status_code:
                parts.append(f"status={status_code}")
            if reason:
                parts.append(f"reason={reason}")
        except Exception:
            pass

        try:
            content = getattr(exc, "content", b"")
            if isinstance(content, bytes):
                content_text = content.decode("utf-8", errors="replace").strip()
            else:
                content_text = str(content).strip()

            if content_text:
                try:
                    parsed = json.loads(content_text)
                    pretty = json.dumps(parsed, ensure_ascii=False)
                    parts.append(f"google_response={pretty}")
                except Exception:
                    parts.append(f"google_response={content_text}")
        except Exception:
            pass

    if not parts:
        parts.append(f"{exc.__class__.__name__}: empty error message")

    return " | ".join(parts)


def _ensure_parent(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)


def _ensure_backup_dir() -> None:
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)


def _normalize_path(value: str) -> Path:
    path = Path(value)
    if not path.is_absolute():
        path = Path.cwd() / path
    return path.resolve()


def _create_backup(path: Path) -> Optional[Path]:
    if not path.exists():
        return None
    _ensure_backup_dir()
    backup_name = f"{path.stem}_{_utc_now_stamp()}{path.suffix}"
    backup_path = BACKUP_DIR / backup_name
    shutil.copy2(path, backup_path)
    return backup_path


def _is_row_like(value: Any) -> bool:
    return isinstance(value, (list, tuple))


def _normalize_rows(rows: Any) -> List[List[Any]]:
    if rows is None:
        return []
    if not isinstance(rows, list):
        raise ValueError("rows must be a list")
    if not rows:
        return []
    if all(_is_row_like(item) for item in rows):
        return [list(item) for item in rows]
    return [list(rows)]


def _normalize_headers(headers: Any) -> List[Any]:
    if headers is None:
        return []
    if not isinstance(headers, list):
        raise ValueError("headers must be a list")
    return list(headers)


def _autofit_worksheet(ws) -> None:
    widths: Dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            value_len = len(str(cell.value))
            col_idx = cell.column
            widths[col_idx] = max(widths.get(col_idx, 0), value_len)

    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(width + 2, 10), 40)


def _format_header(ws) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    _autofit_worksheet(ws)


def _worksheet_summary(ws) -> Dict[str, Any]:
    return {
        "title": ws.title,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
    }


def _inspect_local_excel(path: Path) -> Dict[str, Any]:
    wb = load_workbook(path)
    return {
        "path": str(path),
        "sheets": [_worksheet_summary(ws) for ws in wb.worksheets],
        "active_sheet": wb.active.title if wb.active else None,
    }


def _create_local_excel(payload: Dict[str, Any]) -> Dict[str, Any]:
    output_path = _normalize_path(payload["path"])
    sheet_name = payload.get("sheet_name") or "Sheet1"
    headers = _normalize_headers(payload.get("headers"))
    rows = _normalize_rows(payload.get("rows"))
    apply_header_format = bool(payload.get("format_header", True))

    _ensure_parent(output_path)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    if headers:
        ws.append(headers)
    for row in rows:
        ws.append(row)

    if headers and apply_header_format:
        _format_header(ws)

    wb.save(output_path)

    return {
        "ok": True,
        "action": "create_table",
        "mode": "local_excel",
        "path": str(output_path),
        "sheet": sheet_name,
        "rows_written": len(rows),
        "headers_count": len(headers),
    }


def _append_rows_local_excel(payload: Dict[str, Any]) -> Dict[str, Any]:
    output_path = _normalize_path(payload["path"])
    if not output_path.exists():
        raise FileNotFoundError(f"Workbook not found: {output_path}")

    rows = _normalize_rows(payload.get("rows"))
    sheet_name = payload.get("sheet_name")
    backup_path = _create_backup(output_path)

    wb = load_workbook(output_path)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

    for row in rows:
        ws.append(row)

    _autofit_worksheet(ws)
    wb.save(output_path)

    return {
        "ok": True,
        "action": "append_rows",
        "mode": "local_excel",
        "path": str(output_path),
        "sheet": ws.title,
        "rows_appended": len(rows),
        "backup_path": str(backup_path) if backup_path else None,
    }


def _update_cells_local_excel(payload: Dict[str, Any]) -> Dict[str, Any]:
    output_path = _normalize_path(payload["path"])
    if not output_path.exists():
        raise FileNotFoundError(f"Workbook not found: {output_path}")

    updates = payload.get("updates") or []
    sheet_name = payload.get("sheet_name")
    backup_path = _create_backup(output_path)

    wb = load_workbook(output_path)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

    applied = []
    for item in updates:
        cell_ref = item["cell"]
        value = item.get("value")
        ws[cell_ref] = value
        applied.append({"cell": cell_ref, "value": value})

    _autofit_worksheet(ws)
    wb.save(output_path)

    return {
        "ok": True,
        "action": "update_cells",
        "mode": "local_excel",
        "path": str(output_path),
        "sheet": ws.title,
        "applied_updates": applied,
        "backup_path": str(backup_path) if backup_path else None,
    }


def _format_header_local_excel(payload: Dict[str, Any]) -> Dict[str, Any]:
    output_path = _normalize_path(payload["path"])
    if not output_path.exists():
        raise FileNotFoundError(f"Workbook not found: {output_path}")

    sheet_name = payload.get("sheet_name")
    backup_path = _create_backup(output_path)

    wb = load_workbook(output_path)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

    _format_header(ws)
    wb.save(output_path)

    return {
        "ok": True,
        "action": "format_header",
        "mode": "local_excel",
        "path": str(output_path),
        "sheet": ws.title,
        "backup_path": str(backup_path) if backup_path else None,
    }


def _inspect_local_csv(path: Path) -> Dict[str, Any]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    return {
        "path": str(path),
        "rows": len(rows),
        "columns": len(rows[0]) if rows else 0,
        "preview": rows[:10],
    }


def _create_local_csv(payload: Dict[str, Any]) -> Dict[str, Any]:
    output_path = _normalize_path(payload["path"])
    headers = _normalize_headers(payload.get("headers"))
    rows = _normalize_rows(payload.get("rows"))

    _ensure_parent(output_path)
    with output_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        if headers:
            writer.writerow(headers)
        writer.writerows(rows)

    return {
        "ok": True,
        "action": "create_table",
        "mode": "local_excel",
        "path": str(output_path),
        "rows_written": len(rows),
        "headers_count": len(headers),
    }


def _append_rows_local_csv(payload: Dict[str, Any]) -> Dict[str, Any]:
    output_path = _normalize_path(payload["path"])
    if not output_path.exists():
        raise FileNotFoundError(f"CSV not found: {output_path}")

    rows = _normalize_rows(payload.get("rows"))
    backup_path = _create_backup(output_path)

    with output_path.open("a", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerows(rows)

    return {
        "ok": True,
        "action": "append_rows",
        "mode": "local_excel",
        "path": str(output_path),
        "rows_appended": len(rows),
        "backup_path": str(backup_path) if backup_path else None,
    }


def _get_oauth_credentials_path() -> Path:
    value = os.getenv("GOOGLE_OAUTH_CLIENT_SECRETS", "").strip()
    if not value:
        raise RuntimeError("GOOGLE_OAUTH_CLIENT_SECRETS is not configured")
    path = Path(value)
    if not path.exists():
        raise FileNotFoundError(f"GOOGLE_OAUTH_CLIENT_SECRETS file not found: {path}")
    return path


def _get_oauth_token_path() -> Path:
    value = os.getenv("GOOGLE_OAUTH_TOKEN_JSON", "").strip()
    if value:
        return Path(value)
    return DEFAULT_OAUTH_TOKEN_PATH


def _get_google_mode(payload: Dict[str, Any]) -> str:
    mode = (payload.get("auth_mode") or os.getenv("GOOGLE_AUTH_MODE", "service_account")).strip().lower()
    if mode not in {"service_account", "oauth_user"}:
        raise ValueError(f"Unsupported Google auth_mode: {mode}")
    return mode


def _service_account_credentials(scopes: List[str]):
    creds_path = os.getenv("GOOGLE_SERVICE_ACCOUNT_JSON", "").strip()
    if not creds_path:
        raise RuntimeError("GOOGLE_SERVICE_ACCOUNT_JSON is not configured")
    if not Path(creds_path).exists():
        raise FileNotFoundError(f"GOOGLE_SERVICE_ACCOUNT_JSON file not found: {creds_path}")

    from google.oauth2.service_account import Credentials
    return Credentials.from_service_account_file(creds_path, scopes=scopes)


def _oauth_user_credentials(scopes: List[str], allow_interactive: bool = False):
    from google.auth.transport.requests import Request
    from google.oauth2.credentials import Credentials
    from google_auth_oauthlib.flow import InstalledAppFlow

    token_path = _get_oauth_token_path()
    creds = None

    if token_path.exists():
        creds = Credentials.from_authorized_user_file(str(token_path), scopes)

    if creds and creds.valid:
        return creds

    if creds and creds.expired and creds.refresh_token:
        creds.refresh(Request())
        token_path.parent.mkdir(parents=True, exist_ok=True)
        token_path.write_text(creds.to_json(), encoding="utf-8")
        return creds

    if not allow_interactive:
        raise RuntimeError(
            "OAuth user token is missing or expired and interactive login is disabled. "
            "Run google_sheets/oauth_bootstrap first."
        )

    client_secrets_path = _get_oauth_credentials_path()
    flow = InstalledAppFlow.from_client_secrets_file(str(client_secrets_path), scopes)
    creds = flow.run_local_server(port=0, open_browser=True)
    token_path.parent.mkdir(parents=True, exist_ok=True)
    token_path.write_text(creds.to_json(), encoding="utf-8")
    return creds


def _resolve_google_credentials(payload: Dict[str, Any], allow_interactive: bool = False):
    auth_mode = _get_google_mode(payload)
    if auth_mode == "oauth_user":
        return _oauth_user_credentials(GOOGLE_SCOPES, allow_interactive=allow_interactive)
    return _service_account_credentials(GOOGLE_SCOPES)


def _google_sheets_service(payload: Dict[str, Any], allow_interactive: bool = False):
    creds = _resolve_google_credentials(payload, allow_interactive=allow_interactive)
    from googleapiclient.discovery import build
    return build("sheets", "v4", credentials=creds, cache_discovery=False)


def _google_drive_service(payload: Dict[str, Any], allow_interactive: bool = False):
    creds = _resolve_google_credentials(payload, allow_interactive=allow_interactive)
    from googleapiclient.discovery import build
    return build("drive", "v3", credentials=creds, cache_discovery=False)


def google_probe(payload: Dict[str, Any]) -> Dict[str, Any]:
    auth_mode = _get_google_mode(payload)
    result: Dict[str, Any] = {
        "ok": False,
        "auth_mode": auth_mode,
        "scopes": GOOGLE_SCOPES,
    }

    if auth_mode == "oauth_user":
        client_path = _get_oauth_credentials_path()
        token_path = _get_oauth_token_path()

        result["oauth_client_secrets_path"] = str(client_path)
        result["oauth_client_secrets_exists"] = client_path.exists()
        result["oauth_token_path"] = str(token_path)
        result["oauth_token_exists"] = token_path.exists()

        with client_path.open("r", encoding="utf-8") as f:
            raw = json.load(f)
        result["oauth_client_top_level_keys"] = list(raw.keys())

        creds = _resolve_google_credentials(payload, allow_interactive=False)
        result["credentials_loaded"] = True
        result["credentials_class"] = creds.__class__.__name__
    else:
        creds_path = os.getenv("GOOGLE_SERVICE_ACCOUNT_JSON", "").strip()
        result["env_present"] = bool(creds_path)
        result["env_path"] = creds_path or None
        result["env_path_exists"] = Path(creds_path).exists() if creds_path else False

        if not creds_path:
            raise RuntimeError("GOOGLE_SERVICE_ACCOUNT_JSON is not configured")
        if not Path(creds_path).exists():
            raise FileNotFoundError(f"GOOGLE_SERVICE_ACCOUNT_JSON file not found: {creds_path}")

        with open(creds_path, "r", encoding="utf-8") as f:
            raw = json.load(f)

        result["credential_type"] = raw.get("type")
        result["project_id"] = raw.get("project_id")
        result["client_email"] = raw.get("client_email")

        creds = _resolve_google_credentials(payload, allow_interactive=False)
        result["credentials_loaded"] = True
        result["service_account_email"] = getattr(creds, "service_account_email", None)

    sheets_service = _google_sheets_service(payload, allow_interactive=False)
    result["sheets_service_created"] = True

    drive_service = _google_drive_service(payload, allow_interactive=False)
    result["drive_service_created"] = True

    try:
        about = drive_service.about().get(fields="user").execute()
        result["drive_about_ok"] = True
        result["drive_about_user"] = about.get("user", {})
    except Exception as exc:
        result["drive_about_ok"] = False
        result["drive_about_error"] = _format_exception(exc)

    result["ok"] = True
    return result


def google_oauth_bootstrap(payload: Dict[str, Any]) -> Dict[str, Any]:
    auth_mode = _get_google_mode(payload)
    if auth_mode != "oauth_user":
        raise RuntimeError("google_oauth_bootstrap requires auth_mode=oauth_user")

    creds = _resolve_google_credentials(payload, allow_interactive=True)
    token_path = _get_oauth_token_path()

    drive_service = _google_drive_service(payload, allow_interactive=False)
    about = drive_service.about().get(fields="user").execute()

    return {
        "ok": True,
        "auth_mode": auth_mode,
        "token_path": str(token_path),
        "token_exists": token_path.exists(),
        "authorized_user": about.get("user", {}),
        "scopes": GOOGLE_SCOPES,
        "credentials_class": creds.__class__.__name__,
    }


def _share_drive_file(payload: Dict[str, Any], file_id: str, emails: List[str]) -> None:
    if not emails:
        return

    drive = _google_drive_service(payload, allow_interactive=False)
    for email in emails:
        if not isinstance(email, str) or "@" not in email:
            raise ValueError(f"Invalid share_with email: {email}")

        permission_body = {
            "type": "user",
            "role": "writer",
            "emailAddress": email,
        }
        drive.permissions().create(
            fileId=file_id,
            body=permission_body,
            sendNotificationEmail=False,
        ).execute()


def _google_create_table(payload: Dict[str, Any]) -> Dict[str, Any]:
    title = payload.get("title") or f"Jarvis Sheet {_utc_now_stamp()}"
    sheet_name = payload.get("sheet_name") or "Sheet1"
    headers = _normalize_headers(payload.get("headers"))
    rows = _normalize_rows(payload.get("rows"))
    share_with = payload.get("share_with") or []

    service = _google_sheets_service(payload, allow_interactive=False)
    body = {
        "properties": {"title": title},
        "sheets": [{"properties": {"title": sheet_name}}],
    }
    spreadsheet = service.spreadsheets().create(body=body).execute()
    spreadsheet_id = spreadsheet["spreadsheetId"]
    spreadsheet_url = spreadsheet.get("spreadsheetUrl")

    values: List[List[Any]] = []
    if headers:
        values.append(headers)
    values.extend(rows)

    if values:
        service.spreadsheets().values().update(
            spreadsheetId=spreadsheet_id,
            range=f"{sheet_name}!A1",
            valueInputOption="USER_ENTERED",
            body={"values": values},
        ).execute()

    if headers:
        sheets = spreadsheet.get("sheets", [])
        if not sheets:
            meta = service.spreadsheets().get(spreadsheetId=spreadsheet_id).execute()
            sheets = meta.get("sheets", [])

        if not sheets:
            raise RuntimeError("Unable to resolve sheet metadata after spreadsheet creation")

        sheet_id = sheets[0]["properties"]["sheetId"]
        requests = [{
            "repeatCell": {
                "range": {
                    "sheetId": sheet_id,
                    "startRowIndex": 0,
                    "endRowIndex": 1,
                },
                "cell": {
                    "userEnteredFormat": {
                        "backgroundColor": {"red": 0.12, "green": 0.31, "blue": 0.47},
                        "horizontalAlignment": "CENTER",
                        "textFormat": {
                            "foregroundColor": {"red": 1, "green": 1, "blue": 1},
                            "bold": True
                        },
                    }
                },
                "fields": "userEnteredFormat(backgroundColor,textFormat,horizontalAlignment)"
            }
        }]
        service.spreadsheets().batchUpdate(
            spreadsheetId=spreadsheet_id,
            body={"requests": requests}
        ).execute()

    if share_with:
        _share_drive_file(payload, spreadsheet_id, share_with)

    return {
        "ok": True,
        "action": "create_table",
        "mode": "google_sheets",
        "auth_mode": _get_google_mode(payload),
        "spreadsheet_id": spreadsheet_id,
        "spreadsheet_url": spreadsheet_url,
        "sheet_name": sheet_name,
        "rows_written": len(rows),
        "headers_count": len(headers),
        "shared_with": share_with,
    }


def _google_append_rows(payload: Dict[str, Any]) -> Dict[str, Any]:
    spreadsheet_id = payload["spreadsheet_id"]
    sheet_name = payload.get("sheet_name") or "Sheet1"
    rows = _normalize_rows(payload.get("rows"))

    service = _google_sheets_service(payload, allow_interactive=False)
    response = service.spreadsheets().values().append(
        spreadsheetId=spreadsheet_id,
        range=f"{sheet_name}!A1",
        valueInputOption="USER_ENTERED",
        insertDataOption="INSERT_ROWS",
        body={"values": rows},
    ).execute()

    return {
        "ok": True,
        "action": "append_rows",
        "mode": "google_sheets",
        "auth_mode": _get_google_mode(payload),
        "spreadsheet_id": spreadsheet_id,
        "sheet_name": sheet_name,
        "rows_appended": len(rows),
        "updates": response.get("updates", {}),
    }


def _google_update_cells(payload: Dict[str, Any]) -> Dict[str, Any]:
    spreadsheet_id = payload["spreadsheet_id"]
    updates = payload.get("updates") or []

    service = _google_sheets_service(payload, allow_interactive=False)
    data = []
    for item in updates:
        sheet_name = item.get("sheet_name") or "Sheet1"
        cell = item["cell"]
        value = item.get("value")
        data.append({
            "range": f"{sheet_name}!{cell}",
            "values": [[value]],
        })

    response = service.spreadsheets().values().batchUpdate(
        spreadsheetId=spreadsheet_id,
        body={
            "valueInputOption": "USER_ENTERED",
            "data": data,
        },
    ).execute()

    return {
        "ok": True,
        "action": "update_cells",
        "mode": "google_sheets",
        "auth_mode": _get_google_mode(payload),
        "spreadsheet_id": spreadsheet_id,
        "updated_ranges": response.get("responses", []),
        "applied_count": len(updates),
    }


def _google_inspect(payload: Dict[str, Any]) -> Dict[str, Any]:
    spreadsheet_id = payload["spreadsheet_id"]
    preview_range = payload.get("preview_range") or "Sheet1!A1:Z20"

    service = _google_sheets_service(payload, allow_interactive=False)
    meta = service.spreadsheets().get(spreadsheetId=spreadsheet_id).execute()
    values = service.spreadsheets().values().get(
        spreadsheetId=spreadsheet_id,
        range=preview_range,
    ).execute()

    return {
        "ok": True,
        "action": "inspect",
        "mode": "google_sheets",
        "auth_mode": _get_google_mode(payload),
        "spreadsheet_id": spreadsheet_id,
        "title": meta.get("properties", {}).get("title"),
        "sheets": [s.get("properties", {}).get("title") for s in meta.get("sheets", [])],
        "preview_range": preview_range,
        "preview_values": values.get("values", []),
    }


def execute_spreadsheet_job(payload: Dict[str, Any]) -> Dict[str, Any]:
    mode = (payload.get("mode") or "local_excel").strip().lower()
    action = (payload.get("action") or "inspect").strip().lower()

    if mode not in {"local_excel", "google_sheets"}:
        raise ValueError(f"Unsupported mode: {mode}")

    try:
        if mode == "local_excel":
            path = _normalize_path(payload["path"])
            suffix = path.suffix.lower()

            if action == "inspect":
                if suffix == ".csv":
                    return _inspect_local_csv(path)
                return _inspect_local_excel(path)

            if action == "create_table":
                if suffix == ".csv":
                    return _create_local_csv(payload)
                return _create_local_excel(payload)

            if action == "append_rows":
                if suffix == ".csv":
                    return _append_rows_local_csv(payload)
                return _append_rows_local_excel(payload)

            if action == "update_cells":
                if suffix == ".csv":
                    raise ValueError("update_cells is not supported for CSV. Use .xlsx instead.")
                return _update_cells_local_excel(payload)

            if action == "format_header":
                if suffix == ".csv":
                    raise ValueError("format_header is not supported for CSV. Use .xlsx instead.")
                return _format_header_local_excel(payload)

            raise ValueError(f"Unsupported local_excel action: {action}")

        if mode == "google_sheets":
            if action == "probe":
                return google_probe(payload)
            if action == "oauth_bootstrap":
                return google_oauth_bootstrap(payload)
            if action == "create_table":
                return _google_create_table(payload)
            if action == "append_rows":
                return _google_append_rows(payload)
            if action == "update_cells":
                return _google_update_cells(payload)
            if action == "inspect":
                return _google_inspect(payload)
            raise ValueError(f"Unsupported google_sheets action: {action}")

        raise ValueError(f"Unsupported mode: {mode}")

    except Exception as exc:
        formatted = _format_exception(exc)
        tb = traceback.format_exc()
        _log_error(f"mode={mode} action={action} error={formatted}\n{tb}")
        raise RuntimeError(formatted) from exc
