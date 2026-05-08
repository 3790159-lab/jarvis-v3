from __future__ import annotations

"""File parsers for Jarvis Phase 13.

Each parser returns a dict with standard fields:
  ok: bool
  text: str (full extracted text, max ~10k chars)
  tables: list[list[list[str]]]  (rows of cells)
  metadata: dict (page_count, paragraph_count, sheet_names, etc.)
  error: str  (empty if ok)
"""

import csv
import io
import json
import os
from pathlib import Path
from typing import Any, Dict, List, Optional


def parse_pdf(path: str) -> Dict[str, Any]:
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            pages_text = []
            tables: List[List[List[str]]] = []
            for page in pdf.pages:
                text = page.extract_text() or ""
                pages_text.append(text)
                for tbl in page.extract_tables() or []:
                    tables.append([[str(cell or "") for cell in row] for row in tbl])
            full_text = "\n\n".join(pages_text)
            return {
                "ok": True,
                "text": full_text[:10000],
                "tables": tables,
                "metadata": {"page_count": len(pdf.pages), "table_count": len(tables)},
                "error": "",
            }
    except Exception as e:
        return {"ok": False, "text": "", "tables": [], "metadata": {}, "error": str(e)}


def parse_docx(path: str) -> Dict[str, Any]:
    try:
        from docx import Document
        doc = Document(path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        full_text = "\n".join(paragraphs)
        tables: List[List[List[str]]] = []
        for tbl in doc.tables:
            rows = []
            for row in tbl.rows:
                rows.append([cell.text.strip() for cell in row.cells])
            tables.append(rows)
        return {
            "ok": True,
            "text": full_text[:10000],
            "tables": tables,
            "metadata": {"paragraph_count": len(paragraphs), "table_count": len(tables)},
            "error": "",
        }
    except Exception as e:
        return {"ok": False, "text": "", "tables": [], "metadata": {}, "error": str(e)}


def parse_xlsx(path: str) -> Dict[str, Any]:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        sheets: List[Dict[str, Any]] = []
        all_text_parts = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows: List[List[str]] = []
            for row in ws.iter_rows(values_only=True):
                rows.append([str(cell) if cell is not None else "" for cell in row])
            sheets.append({"name": sheet_name, "rows": rows})
            if rows:
                all_text_parts.append(f"Sheet: {sheet_name}")
                all_text_parts.extend(["\t".join(r) for r in rows[:50]])
        return {
            "ok": True,
            "text": "\n".join(all_text_parts)[:10000],
            "tables": [s["rows"] for s in sheets],
            "metadata": {"sheet_names": [s["name"] for s in sheets], "sheet_count": len(sheets)},
            "error": "",
        }
    except Exception as e:
        return {"ok": False, "text": "", "tables": [], "metadata": {}, "error": str(e)}


def parse_csv_file(path: str) -> Dict[str, Any]:
    try:
        rows: List[List[str]] = []
        with open(path, encoding="utf-8-sig", newline="") as f:
            reader = csv.reader(f)
            for row in reader:
                rows.append(row)
        text = "\n".join([",".join(r) for r in rows[:200]])
        return {
            "ok": True,
            "text": text[:10000],
            "tables": [rows],
            "metadata": {"row_count": len(rows), "col_count": len(rows[0]) if rows else 0},
            "error": "",
        }
    except Exception as e:
        return {"ok": False, "text": "", "tables": [], "metadata": {}, "error": str(e)}


def parse_image(path: str) -> Dict[str, Any]:
    """Extract text from image via OCR (pytesseract) or return placeholder."""
    try:
        from PIL import Image
        img = Image.open(path)
        width, height = img.size
        mode = img.mode
        try:
            import pytesseract
            text = pytesseract.image_to_string(img, lang="rus+eng")
        except Exception:
            text = ""  # OCR not available
        return {
            "ok": True,
            "text": text[:5000],
            "tables": [],
            "metadata": {"width": width, "height": height, "mode": mode, "ocr_available": bool(text)},
            "error": "",
        }
    except Exception as e:
        return {"ok": False, "text": "", "tables": [], "metadata": {}, "error": str(e)}


def parse_text_file(path: str) -> Dict[str, Any]:
    try:
        text = Path(path).read_text(encoding="utf-8", errors="replace")
        lines = text.splitlines()
        return {
            "ok": True,
            "text": text[:10000],
            "tables": [],
            "metadata": {"line_count": len(lines), "char_count": len(text)},
            "error": "",
        }
    except Exception as e:
        return {"ok": False, "text": "", "tables": [], "metadata": {}, "error": str(e)}


# Extension → parser mapping
_EXTENSION_MAP = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".doc": parse_docx,
    ".xlsx": parse_xlsx,
    ".xls": parse_xlsx,
    ".csv": parse_csv_file,
    ".txt": parse_text_file,
    ".md": parse_text_file,
    ".json": parse_text_file,
    ".png": parse_image,
    ".jpg": parse_image,
    ".jpeg": parse_image,
    ".webp": parse_image,
}

_MIME_MAP = {
    "application/pdf": parse_pdf,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": parse_docx,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": parse_xlsx,
    "text/csv": parse_csv_file,
    "text/plain": parse_text_file,
    "image/png": parse_image,
    "image/jpeg": parse_image,
    "image/webp": parse_image,
}


def parse_file(path: str, mime_type: str = "") -> Dict[str, Any]:
    """Auto-detect parser from extension or mime_type and parse the file."""
    ext = Path(path).suffix.lower()
    parser = _EXTENSION_MAP.get(ext) or _MIME_MAP.get(mime_type)
    if parser is None:
        return {
            "ok": False, "text": "", "tables": [], "metadata": {},
            "error": f"Unsupported file type: ext={ext}, mime={mime_type}",
        }
    return parser(path)


def summarize_parse_result(result: Dict[str, Any], filename: str = "") -> str:
    """Human-readable one-line summary of a parse result."""
    if not result.get("ok"):
        return f"⚠️ Не удалось прочитать файл: {result.get('error', 'unknown error')}"
    meta = result.get("metadata") or {}
    parts = []
    if "page_count" in meta:
        parts.append(f"{meta['page_count']} стр.")
    if "paragraph_count" in meta:
        parts.append(f"{meta['paragraph_count']} абзацев")
    if "row_count" in meta:
        parts.append(f"{meta['row_count']} строк")
    if "sheet_names" in meta:
        parts.append(f"листы: {', '.join(meta['sheet_names'][:3])}")
    if "width" in meta:
        parts.append(f"{meta['width']}×{meta['height']} px")
    text_len = len(result.get("text") or "")
    parts.append(f"~{text_len} символов")
    detail = ", ".join(parts) if parts else "пустой файл"
    name = filename or "файл"
    return f"📄 {name} ({detail})"
